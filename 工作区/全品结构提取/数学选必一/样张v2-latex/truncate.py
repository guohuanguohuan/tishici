# -*- coding: utf-8 -*-
"""复制源 docx 并截断：从首个「样式=标题3(Heading3) 且文本以 1.1.2 起首」的段落起删至文末（保留 sectPr）。
另登记：该截断段之前紧邻的 1.1.2 节名锚（1pt 白字隐藏段）一并删除，因任务枚举的保留范围是「章首＋1.1 节标题＋1.1.1 全节」。"""
import shutil
from docx import Document
from docx.oxml.ns import qn

SRC = r"C:\提示词\高中数学\高中数学同步\人教B版选必1 第1章 空间向量与立体几何（上）·讲练件（61题）.docx"
OUT = r"C:\提示词\工作区\全品结构提取\数学选必一\样张v2-latex\1.1.1-源节.docx"

shutil.copyfile(SRC, OUT)
doc = Document(OUT)

def is_heading3(p):
    return p.style.style_id in ('Heading3', '3') or p.style.name in ('标题3', '标题 3', 'Heading 3', 'Heading3')

def text(p):
    return ''.join(node.text or '' for node in p._element.iter() if node.tag in (qn('w:t'), qn('m:t')))

paras = doc.paragraphs
cut = None
for p in paras:
    if is_heading3(p) and text(p).strip().startswith('1.1.2'):
        cut = p
        break
assert cut is not None, '未找到 1.1.2 标题3 段'
print('截断段:', repr(text(cut)[:30]), 'style=', cut.style.style_id)

# 先（在 cut 仍挂在 body 上时）定位紧邻的 1.1.2 节名锚——隐藏锚不属于保留范围枚举
anchor = None
prev = cut._element.getprevious()
while prev is not None and prev.tag == qn('w:p'):
    t = ''.join(n.text or '' for n in prev.iter() if n.tag == qn('w:t'))
    if t.strip():
        if t.strip().startswith('1.1.2'):
            anchor = prev
        break
    prev = prev.getprevious()

body = doc.element.body
sectPr = body.find(qn('w:sectPr'))
removed = 0
hit = False
for el in list(body):
    if el is cut._element:
        hit = True
    if hit and el is not sectPr:
        body.remove(el)
        removed += 1
print('删除 body 子元素:', removed)

if anchor is not None:
    body.remove(anchor)
    print('另删 1.1.2 节名锚')

doc.save(OUT)
print('已保存:', OUT)

# 校验
d2 = Document(OUT)
n_tbl = len(d2.tables)
n_par = len(d2.paragraphs)
print(f'校验: 段落 {n_par}, 表格 {n_tbl}, 末段文本: {d2.paragraphs[-1].text[:40]!r}')
