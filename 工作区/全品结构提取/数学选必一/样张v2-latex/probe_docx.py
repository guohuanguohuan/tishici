# -*- coding: utf-8 -*-
"""探查源 docx 结构：样式名、1.1.x 标题、表/图/段落数量分布"""
import re
from docx import Document

SRC = r"C:\提示词\高中数学\高中数学同步\人教B版选必1 第1章 空间向量与立体几何（上）·讲练件（61题）.docx"
doc = Document(SRC)
body = doc.element.body

ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
kids = list(body)
print("body 直接子元素总数:", len(kids))
from collections import Counter
tags = Counter(k.tag.split('}')[1] for k in kids)
print("子元素标签分布:", dict(tags))

# 标题段落扫描
print("\n--- 所有含 '1.1' 起首文本的段落 ---")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if re.match(r"^1\.1", t):
        st = p.style
        print(f"para#{i} style_id={st.style_id!r} name={st.name!r} text={t[:40]!r}")

# 样式名普查（前60段）
print("\n--- 前 60 段样式 ---")
for i, p in enumerate(doc.paragraphs[:60]):
    t = p.text.strip()[:24]
    print(f"#{i} [{p.style.style_id}/{p.style.name}] {t!r}")
