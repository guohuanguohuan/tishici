# -*- coding: utf-8 -*-
"""图债工位：按源指向批量提取候选内嵌图（对象层，只读源件）。
输出 源证/提取原图/<债键>_<序>.<ext>＋清单 源证/提取清单.txt
富矿件按题号 #N 定位：段落文本命中题号标记或题面签名后取邻域图。"""
import io, os, re
from docx import Document
from docx.oxml.ns import qn

SRC3 = r"C:\提示词\高中数学\参考\组卷网\【新课标 新探索】大单元作业设计\人教A版选择性必修1\第3章 圆锥曲线的方程"
SRC2 = r"C:\提示词\高中数学\高中数学同步"
OUT = r"C:\提示词\工作区\_tmpM3图债0914\源证\提取原图"
os.makedirs(OUT, exist_ok=True)

VOL = {
    "卷①": "人教B版选必1 第2章 平面解析几何（2.1—2.3.3）·讲练件（92题）.docx",
    "卷②": "人教B版选必1 第2章 平面解析几何（2.3.4—2.5.2）·讲练件（90题）.docx",
    "卷③": "人教B版选必1 第2章 平面解析几何（2.6.1—2.7.2）·讲练件（68题）.docx",
    "卷④": "人教B版选必1 第2章 平面解析几何（2.8）·讲练件（89题）.docx",
}
VML = "{urn:schemas-microsoft-com:vml}imagedata"

def para_images(p):
    out = []
    for blip in p._p.findall(".//" + qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if rid:
            out.append(rid)
    for imd in p._p.findall(".//" + VML):
        rid = imd.get(qn("r:id"))
        if rid:
            out.append(rid)
    return out

def save(doc, rid, name):
    try:
        part = doc.part.related_parts[rid]
    except KeyError:
        return f"{name}: rId {rid} 不在关系表"
    ext = str(part.partname).rsplit(".", 1)[-1].lower()
    fp = os.path.join(OUT, f"{name}.{ext}")
    with open(fp, "wb") as f:
        f.write(part.blob)
    return f"{os.path.basename(fp)}  {len(part.blob)}B"

log = []

def dump_vol(vkey, debt, prange):
    doc = Document(os.path.join(SRC2, VOL[vkey]))
    paras = doc.paragraphs
    for i in prange:
        for k, rid in enumerate(para_images(paras[i])):
            log.append(save(doc, rid, f"{debt}__{vkey}段{i}_{k}"))

# —— 四卷定点段（probe 实测）——
dump_vol("卷①", "02-E1", [54, 55, 56, 57, 58, 59])
dump_vol("卷②", "10-G2", [222, 223, 224, 225, 226])
dump_vol("卷②", "12-例15", [1281, 1282, 1283, 1284])
dump_vol("卷③", "14-变式34", [422, 423, 424, 425])
dump_vol("卷③", "14-例39", [555, 556, 557, 558])
dump_vol("卷③", "16-变式39", [1050, 1051, 1052, 1053])
dump_vol("卷④", "17-拓01", [659, 660, 661, 662, 663])

# —— 富矿件：题号定位 ——
MINE = [
    ("13-变式16", "13 双曲线的综合问题（共25题）.docx", 25, ["猫捉老鼠", "人工智能大会"]),
    ("14-例46", "13 双曲线的综合问题（共25题）.docx", 22, ["青花瓷"]),
    ("14-变式47", "13 双曲线的综合问题（共25题）.docx", 24, ["唐山", "钓鱼岛"]),
    ("14-变式49", "10 双曲线的几何性质（共48题）.docx", 6, ["篮球"]),
    ("14-变式53", "10 双曲线的几何性质（共48题）.docx", 43, ["大教堂", "南非"]),
    ("15-例23", "18 抛物线的综合问题（共28题）.docx", 26, ["隧道", "集装箱"]),
    ("18-12", "19 圆锥曲线之间的综合问题（共15题）.docx", 14, ["内切圆", "ty+1", "t y + 1"]),
    ("19-变式10", "19 圆锥曲线之间的综合问题（共15题）.docx", 8, ["公共焦点", "离心率之积"]),
    ("19-例16", "19 圆锥曲线之间的综合问题（共15题）.docx", 15, ["交曲线", "中点", "准线"]),
    ("18-变式15", "21 章节综合测试-圆锥曲线的方程（共22题）.docx", 21, ["上半椭圆", "连接而成"]),
]

for debt, fn, qn_no, sigs in MINE:
    path = os.path.join(SRC3, fn)
    doc = Document(path)
    paras = doc.paragraphs
    texts = [p.text for p in paras]
    # 题号标记（「#N．」或「N．」行首）与签名双轨定位
    cand = []
    pat = re.compile(rf"(?:^|[\s＞>　])#{qn_no}[．.、]|(?:^|\n){qn_no}\s*[．.、]")
    for i, t in enumerate(texts):
        head = t.strip()[:12]
        if re.match(rf"^\D{{0,4}}{qn_no}\s*[．.、]", head):
            cand.append(i)
    for sig in sigs:
        cand += [i for i, t in enumerate(texts) if sig in t]
    cand = sorted(set(cand))
    log.append(f"### {debt} 件「{fn}」#{qn_no} 候选段 {cand[:12]}")
    if not cand:
        log.append("  !! 未定位")
        continue
    i0 = cand[0]
    got = 0
    for j in range(max(0, i0 - 2), min(len(paras), i0 + 14)):
        rids = para_images(paras[j])
        for rid in rids:
            log.append("  " + save(doc, rid, f"{debt}__件段{j}_{got}"))
            got += 1
    if not got:
        # 邻域无图：再扫远一点并记录
        far = [(j, para_images(paras[j])) for j in range(max(0, i0-15), min(len(paras), i0+30)) if para_images(paras[j])]
        log.append(f"  邻域±(2,14)无图；远邻域: {far[:6] if far else '无'}")

with io.open(r"C:\提示词\工作区\_tmpM3图债0914\源证\提取清单.txt", "w", encoding="utf-8") as f:
    f.write("\n".join(log))
print("\n".join(log))
